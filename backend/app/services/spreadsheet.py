"""Bounded, deterministic CSV/XLSX table analysis."""

from __future__ import annotations

import csv
import math
import re
from collections import defaultdict
from datetime import date
from typing import TYPE_CHECKING, Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

if TYPE_CHECKING:
    from pathlib import Path

MAX_ROWS = 50_000
MAX_COLUMNS = 100
MAX_RESULT_ROWS = 100
MAX_EVIDENCE_ROWS = 20
MAX_RESULT_CELL_CHARS = 500


class SpreadsheetFilter(BaseModel):
    model_config = ConfigDict(extra="forbid")

    column: str = Field(min_length=1, max_length=256)
    operator: Literal["eq", "ne", "gt", "gte", "lt", "lte", "contains"]
    value: str = Field(default="", max_length=2_000)
    compare_column: str | None = Field(default=None, max_length=256)

    @model_validator(mode="after")
    def validate_comparison(self) -> SpreadsheetFilter:
        if self.compare_column and self.operator == "contains":
            raise ValueError("contains does not support comparing two columns")
        return self


class SpreadsheetPlan(BaseModel):
    """Allowlisted operations the model may request; never executable code."""

    model_config = ConfigDict(
        extra="forbid",
        json_schema_extra={
            "required": ["filters", "group_by", "aggregate", "supported", "column"]
        },
    )

    filters: list[SpreadsheetFilter] = Field(default_factory=list, max_length=20)
    group_by: list[str] = Field(default_factory=list, max_length=10)
    aggregate: Literal["sum", "avg", "count", "min", "max"] | None = None
    supported: bool = True
    column: str | None = Field(default=None, max_length=256)
    sort_by: str | None = Field(default=None, max_length=256)
    descending: bool = False
    limit: int = Field(default=20, ge=1, le=MAX_RESULT_ROWS)

    @model_validator(mode="after")
    def validate_aggregate(self) -> SpreadsheetPlan:
        if self.aggregate is None and (self.group_by or self.column):
            raise ValueError("group_by and column require an aggregate")
        if self.aggregate in {"sum", "avg", "min", "max"} and not self.column:
            raise ValueError(f"{self.aggregate} requires a column")
        if self.aggregate == "count" and self.column is not None:
            raise ValueError("count does not accept a column")
        return self


class SpreadsheetError(ValueError):
    """Invalid or unsupported uploaded table."""


def load_spreadsheet(path: Path, file_format: str) -> dict[str, object]:
    """Load CSV or the first XLSX sheet with source row numbers."""
    workbook = None
    stream = None
    if file_format == "csv":
        stream = path.open(newline="", encoding="utf-8-sig")
        try:
            sample = stream.read(8192)
            stream.seek(0)
            dialect = csv.Sniffer().sniff(sample) if sample else csv.excel
        except csv.Error:
            dialect = csv.excel
        except UnicodeError as exc:
            stream.close()
            raise SpreadsheetError(f"Could not read CSV: {exc}") from exc
        rows = csv.reader(stream, dialect)
        try:
            header_row = next(rows, None)
        except (csv.Error, UnicodeError) as exc:
            stream.close()
            raise SpreadsheetError(f"Could not read CSV: {exc}") from exc
        sheet_names: list[str] = []
        sheet_name = None
    elif file_format == "xlsx":
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet_names = workbook.sheetnames
            sheet_name = sheet_names[0] if sheet_names else None
            rows = (
                iter(workbook[sheet_name].iter_rows(values_only=True)) if sheet_name else iter(())
            )
            header_row = next(rows, None)
        except Exception as exc:
            if workbook is not None:
                workbook.close()
            raise SpreadsheetError(f"Could not read XLSX: {exc}") from exc
    else:
        raise SpreadsheetError("Only CSV and XLSX files are supported.")

    def close_sources() -> None:
        if workbook is not None:
            workbook.close()
        if stream is not None:
            stream.close()

    if header_row is None:
        close_sources()
        raise SpreadsheetError("The spreadsheet is empty.")
    headers = [str(value).strip() if value is not None else "" for value in header_row]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not header for header in headers):
        close_sources()
        raise SpreadsheetError("Every column must have a non-empty header.")
    if len(set(headers)) != len(headers):
        close_sources()
        raise SpreadsheetError("Column headers must be unique.")
    if len(headers) > MAX_COLUMNS:
        close_sources()
        raise SpreadsheetError(f"Spreadsheets are limited to {MAX_COLUMNS} columns.")
    if any(len(header) > 128 for header in headers):
        close_sources()
        raise SpreadsheetError("Column headers are limited to 128 characters.")

    records: list[dict[str, str]] = []
    source_rows: list[int] = []
    try:
        for row_number, row in enumerate(rows, start=2):
            if row_number - 2 >= MAX_ROWS:
                raise SpreadsheetError(f"Spreadsheets are limited to {MAX_ROWS:,} data rows.")
            values = list(row)
            if any(value is not None and str(value).strip() for value in values[len(headers) :]):
                raise SpreadsheetError("A data row contains values beyond the header columns.")
            values = values[: len(headers)]
            values.extend([None] * (len(headers) - len(values)))
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            records.append(
                {
                    header: "" if value is None else str(value).strip()
                    for header, value in zip(headers, values, strict=True)
                }
            )
            source_rows.append(row_number)
    except (csv.Error, UnicodeError) as exc:
        raise SpreadsheetError(f"Could not read spreadsheet rows: {exc}") from exc
    except SpreadsheetError:
        raise
    except Exception as exc:
        if workbook is None:
            raise
        raise SpreadsheetError(f"Could not read XLSX rows: {exc}") from exc
    finally:
        close_sources()

    types = {header: _column_type([record[header] for record in records]) for header in headers}
    return {
        "headers": headers,
        "rows": records,
        "source_rows": source_rows,
        "column_types": types,
        "sheet": sheet_name,
        "available_sheets": sheet_names,
    }


def execute_plan(table: dict[str, object], plan: SpreadsheetPlan) -> dict[str, object]:
    headers = table["headers"]
    rows = table["rows"]
    source_rows = table["source_rows"]
    types = table["column_types"]
    assert isinstance(headers, list) and isinstance(rows, list)
    assert isinstance(source_rows, list) and isinstance(types, dict)
    for column in [
        column for item in plan.filters for column in (item.column, item.compare_column) if column
    ] + plan.group_by:
        if column not in headers:
            raise SpreadsheetError(f"Unknown column: {column}")
    for item in plan.filters:
        if item.compare_column and types[item.column] != types[item.compare_column]:
            raise SpreadsheetError("Compared columns must have the same detected type.")
    if plan.column and plan.column not in headers:
        raise SpreadsheetError(f"Unknown column: {plan.column}")
    aggregate_header = (
        ("count" if plan.aggregate == "count" else f"{plan.aggregate}_{plan.column}")
        if plan.aggregate
        else None
    )
    if plan.sort_by and plan.sort_by not in headers and plan.sort_by != aggregate_header:
        raise SpreadsheetError(f"Unknown sort column: {plan.sort_by}")
    if plan.aggregate in {"sum", "avg", "min", "max"} and types[plan.column] != "number":
        raise SpreadsheetError(f"{plan.aggregate} requires a numeric column.")

    selected: list[tuple[dict[str, str], int]] = []
    for row, source_row in zip(rows, source_rows, strict=True):
        if all(
            _matches(
                row[item.column],
                item,
                types[item.column],
                row[item.compare_column] if item.compare_column else None,
            )
            for item in plan.filters
        ):
            selected.append((row, source_row))

    if plan.aggregate is None:
        result_headers = list(headers)
        result_rows = [dict(row) for row, _ in selected]
        evidence = [[source] for _, source in selected]
        evidence_truncated = False
    else:
        groups: dict[tuple[str, ...], list[tuple[dict[str, str], int]]] = defaultdict(list)
        for row, source in selected:
            groups[tuple(row[column] for column in plan.group_by)].append((row, source))
        if not plan.group_by:
            groups[()] = selected
        result_headers = [
            *plan.group_by,
            "count" if plan.aggregate == "count" else f"{plan.aggregate}_{plan.column}",
        ]
        result_rows = []
        evidence = []
        evidence_truncated = False
        for key, members in groups.items():
            values = (
                [_as_number(row[plan.column]) for row, _ in members if row[plan.column] != ""]
                if plan.column
                else []
            )
            values = [value for value in values if value is not None]
            if plan.aggregate == "count":
                value: int | float = len(members)
            elif not values:
                value = ""
            elif plan.aggregate == "sum":
                value = sum(values)
            elif plan.aggregate == "avg":
                value = sum(values) / len(values)
            elif plan.aggregate == "min":
                value = min(values)
            else:
                value = max(values)
            result_rows.append(
                {**dict(zip(plan.group_by, key, strict=True)), result_headers[-1]: value}
            )
            evidence.append([source for _, source in members][:MAX_EVIDENCE_ROWS])
            evidence_truncated |= len(members) > MAX_EVIDENCE_ROWS

    if plan.sort_by:
        present = [
            i for i, row in enumerate(result_rows) if row.get(plan.sort_by) not in (None, "")
        ]
        missing = [i for i, row in enumerate(result_rows) if row.get(plan.sort_by) in (None, "")]
        present.sort(
            key=lambda index: _sort_value(result_rows[index].get(plan.sort_by)),
            reverse=plan.descending,
        )
        order = present + missing
        result_rows = [result_rows[index] for index in order]
        evidence = [evidence[index] for index in order]

    total_results = len(result_rows)
    result_rows = result_rows[: plan.limit]
    evidence = evidence[: plan.limit]
    cells_truncated = False
    for row in result_rows:
        for column, value in row.items():
            if isinstance(value, str) and len(value) > MAX_RESULT_CELL_CHARS:
                row[column] = value[:MAX_RESULT_CELL_CHARS]
                cells_truncated = True
    return {
        "headers": result_headers,
        "rows": result_rows,
        "source_rows": evidence,
        "matched_rows": len(selected),
        "result_count": total_results,
        "evidence_truncated": evidence_truncated,
        "cells_truncated": cells_truncated,
    }


def _column_type(values: list[str]) -> str:
    nonempty = [value for value in values if value]
    if nonempty and all(_as_number(value) is not None for value in nonempty):
        return "number"
    if nonempty and all(_as_date(value) is not None for value in nonempty):
        return "date"
    return "text"


def _as_number(value: str) -> float | None:
    if "," in value and not re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+(\.\d+)?", value):
        return None
    try:
        number = float(value.replace(",", ""))
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _as_date(value: str) -> date | None:
    try:
        return date.fromisoformat(value[:10])
    except ValueError:
        return None


def _matches(
    value: str, condition: SpreadsheetFilter, column_type: str, compare_value: str | None = None
) -> bool:
    expected = compare_value if compare_value is not None else condition.value
    if condition.operator == "contains" and compare_value is None:
        return expected.casefold() in value.casefold()
    if compare_value is not None and (not value or not expected):
        return False
    if column_type == "number":
        left, right = _as_number(value), _as_number(expected)
        if left is None or right is None:
            return False
    elif column_type == "date":
        left, right = _as_date(value), _as_date(expected)
        if left is None or right is None:
            return False
    else:
        left, right = value.casefold(), expected.casefold()
    return {
        "eq": left == right,
        "ne": left != right,
        "gt": left > right,
        "gte": left >= right,
        "lt": left < right,
        "lte": left <= right,
    }[condition.operator]


def _sort_value(value: object) -> tuple[int, object]:
    if isinstance(value, (int, float)):
        return 0, value
    number = _as_number(str(value)) if value is not None else None
    if number is not None:
        return 0, number
    parsed_date = _as_date(str(value)) if value is not None else None
    if parsed_date is not None:
        return 1, parsed_date
    return 2, str(value).casefold()
