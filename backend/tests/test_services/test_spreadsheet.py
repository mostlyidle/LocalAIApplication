from __future__ import annotations

import xml.etree.ElementTree as ET
from typing import TYPE_CHECKING
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from app.services.spreadsheet import (
    SpreadsheetError,
    SpreadsheetPlan,
    execute_plan,
    load_spreadsheet,
)

if TYPE_CHECKING:
    from pathlib import Path


def test_csv_filter_and_numeric_aggregate_keep_source_rows(tmp_path: Path) -> None:
    path = tmp_path / "sales.csv"
    path.write_text("Team,Revenue\nWest,10\nEast,25\nWest,15\n", encoding="utf-8")
    table = load_spreadsheet(path, "csv")
    result = execute_plan(
        table,
        SpreadsheetPlan(
            filters=[{"column": "Revenue", "operator": "gt", "value": "10"}],
            group_by=["Team"],
            aggregate="sum",
            column="Revenue",
            sort_by="sum_Revenue",
            descending=True,
        ),
    )
    assert result["rows"] == [
        {"Team": "East", "sum_Revenue": 25.0},
        {"Team": "West", "sum_Revenue": 15.0},
    ]
    assert result["source_rows"] == [[3], [4]]
    assert table["column_types"] == {"Team": "text", "Revenue": "number"}


def test_xlsx_uses_first_sheet_and_reports_available_sheets(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    sheet.append(["Name", "Amount"])
    sheet.append(["A", 12])
    workbook.create_sheet("Notes").append(["Other"])
    workbook.save(path)

    table = load_spreadsheet(path, "xlsx")
    assert table["sheet"] == "Orders"
    assert table["available_sheets"] == ["Orders", "Notes"]
    assert table["rows"] == [{"Name": "A", "Amount": "12"}]


def test_xlsx_formula_without_cached_result_is_not_recalculated(tmp_path: Path) -> None:
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Team", "Revenue"])
    sheet.append(["West", 10])
    sheet.append(["West", 15])
    sheet.append(["Total", "=SUM(B2:B3)"])
    workbook.save(path)

    table = load_spreadsheet(path, "xlsx")

    assert table["rows"][-1] == {"Team": "Total", "Revenue": ""}


def test_xlsx_formula_uses_saved_cache_instead_of_recalculating(tmp_path: Path) -> None:
    path = tmp_path / "cached-formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Team", "Revenue"])
    sheet.append(["West", 10])
    sheet.append(["West", 15])
    sheet.append(["Total", "=SUM(B2:B3)"])
    workbook.save(path)

    with ZipFile(path) as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as updated:
        for item, content in entries:
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                cached_value = root.find(f".//{namespace}c[@r='B4']/{namespace}v")
                assert cached_value is not None
                cached_value.text = "999"
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            updated.writestr(item, content)

    table = load_spreadsheet(path, "xlsx")

    assert table["rows"][-1] == {"Team": "Total", "Revenue": "999"}


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ("a,a\n1,2\n", "unique"),
        (",b\n1,2\n", "non-empty"),
    ],
)
def test_csv_rejects_ambiguous_headers(tmp_path: Path, content: str, message: str) -> None:
    path = tmp_path / "bad.csv"
    path.write_text(content, encoding="utf-8")
    with pytest.raises(SpreadsheetError, match=message):
        load_spreadsheet(path, "csv")


def test_csv_rejects_nonempty_cells_beyond_header_columns(tmp_path: Path) -> None:
    path = tmp_path / "wide-row.csv"
    path.write_text("Name\nA,unexpected\n", encoding="utf-8")
    with pytest.raises(SpreadsheetError, match="beyond the header"):
        load_spreadsheet(path, "csv")


def test_plan_rejects_unknown_columns_and_non_numeric_aggregates(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("Name\nA\n", encoding="utf-8")
    table = load_spreadsheet(path, "csv")
    with pytest.raises(SpreadsheetError, match="Unknown column"):
        execute_plan(
            table, SpreadsheetPlan(filters=[{"column": "Missing", "operator": "eq", "value": "x"}])
        )
    with pytest.raises(SpreadsheetError, match="numeric"):
        execute_plan(table, SpreadsheetPlan(aggregate="sum", column="Name"))


def test_numeric_aggregates_support_thousands_separators(tmp_path: Path) -> None:
    path = tmp_path / "amounts.csv"
    path.write_text('Amount\n"1,200"\n"2,300"\n', encoding="utf-8")
    table = load_spreadsheet(path, "csv")
    result = execute_plan(table, SpreadsheetPlan(aggregate="sum", column="Amount"))
    assert result["rows"] == [{"sum_Amount": 3500.0}]


def test_invalid_thousands_grouping_is_not_coerced_to_a_number(tmp_path: Path) -> None:
    path = tmp_path / "amounts.csv"
    path.write_text('Amount\n"12,34"\n', encoding="utf-8")
    table = load_spreadsheet(path, "csv")
    assert table["column_types"] == {"Amount": "text"}


def test_filter_can_compare_two_numeric_columns(tmp_path: Path) -> None:
    path = tmp_path / "budget.csv"
    path.write_text("Actual,Approved\n12,10\n8,10\n", encoding="utf-8")
    table = load_spreadsheet(path, "csv")
    result = execute_plan(
        table,
        SpreadsheetPlan(
            filters=[
                {
                    "column": "Actual",
                    "operator": "gt",
                    "compare_column": "Approved",
                }
            ]
        ),
    )
    assert result["rows"] == [{"Actual": "12", "Approved": "10"}]
    assert result["source_rows"] == [[2]]


def test_loader_stops_at_configured_row_limit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import app.services.spreadsheet as spreadsheet

    path = tmp_path / "large.csv"
    path.write_text("Value\n1\n2\n3\n", encoding="utf-8")
    monkeypatch.setattr(spreadsheet, "MAX_ROWS", 2)
    with pytest.raises(SpreadsheetError, match="2 data rows"):
        load_spreadsheet(path, "csv")
